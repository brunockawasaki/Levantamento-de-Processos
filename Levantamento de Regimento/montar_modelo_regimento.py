# -*- coding: utf-8 -*-
"""Gera 'RE.GG.PRO.001 - MODELO REGIMENTO com marcadores.docx' a partir do template oficial de Regimento.

Uso: python montar_modelo_regimento.py <template_RE.docx> <saida_com_marcadores.docx>

Marcadores colocados (so no corpo, igual aos modelos de POP e IT):
  {{TITULO}}                       titulo da capa (caixa de texto e copia de compatibilidade)
  {{OBJETIVO}}                     texto da secao 1
  {{APLICACAO}}                    texto da secao 2
  {{AREA_ITEM}}                    um paragrafo com marcador (bullet) por area (secao 3)
  {{GOV_CONVOCACAO}} ... {{GOV_ATA}}  as 5 celulas de DEFINICAO da tabela de Governanca (secao 4)
  {{AREA}} / {{AREA_RESP}}         linha repetida da tabela de Papeis (secao 5)
  {{INPUT}} / {{INPUT_DESC}} / {{INPUT_RESP}}      linha repetida da tabela de Entradas (secao 6)
  {{OUTPUT}} / {{OUTPUT_DESC}} / {{OUTPUT_RESP}}   linha repetida da tabela de Saidas (secao 7)
  {{REFERENCIA_ITEM}}              um paragrafo com marcador (bullet) por referencia (secao 8)
O Controle de Emissao, a capa (codigo, data, aprovador...), o cabecalho e o rodape ficam como no
template (a equipe de Processos preenche).
"""
import sys
import docx
from docx.oxml.ns import qn
from docx.oxml import parse_xml

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
RPR = ('<w:rPr><w:rFonts w:ascii="Arial" w:eastAsia="Calibri" w:hAnsi="Arial" w:cs="Arial"/>'
       '<w:color w:val="000000" w:themeColor="text1"/><w:sz w:val="20"/><w:szCs w:val="20"/>'
       '<w:lang w:eastAsia="pt-BR"/></w:rPr>')


def ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def set_single_run(p_el, text):
    """Deixa o paragrafo com 1 so run (mantendo o rPr do primeiro) com o texto dado."""
    runs = p_el.findall(qn("w:r"))
    first = runs[0]
    for r in runs[1:]:
        p_el.remove(r)
    for child in list(first):
        if child.tag != qn("w:rPr"):
            first.remove(child)
    t = first.makeelement(qn("w:t"), {})
    t.text = text
    first.append(t)


def cell_token(cell, token):
    """Troca o conteudo da celula por um unico run com o marcador (Arial 10)."""
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    p = paras[0]._p
    for r in p.findall(qn("w:r")):
        p.remove(r)
    p.append(parse_xml('<w:r %s>%s<w:t>%s</w:t></w:r>' % (W, RPR, token)))


def bullet_paragraph(token):
    return parse_xml(
        '<w:p %s><w:pPr><w:pStyle w:val="PargrafodaLista"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr>'
        '<w:spacing w:after="0" w:line="360" w:lineRule="auto"/><w:contextualSpacing w:val="0"/><w:jc w:val="both"/>%s</w:pPr>'
        '<w:r>%s<w:t>%s</w:t></w:r></w:p>' % (W, RPR, RPR, token))


def main(src, dst):
    d = docx.Document(src)
    body_ps = list(d.paragraphs)

    def find(startswith):
        for p in body_ps:
            if ptext(p._p).strip().startswith(startswith):
                return p
        raise SystemExit("Nao achei o paragrafo: " + startswith)

    # 1, 2: textos simples (o texto de instrucao do template vira marcador)
    set_single_run(find("Definir o objetivo")._p, "{{OBJETIVO}}")
    set_single_run(find("Descrever a Empresa")._p, "{{APLICACAO}}")

    # 3 e 8: listas com marcador (bullet), um paragrafo por item
    for inicio, token in (("Mencionar quais", "{{AREA_ITEM}}"), ("Relacionar as normas", "{{REFERENCIA_ITEM}}")):
        alvo = find(inicio)
        alvo._p.addprevious(bullet_paragraph(token))
        alvo._p.getparent().remove(alvo._p)

    # 4: Governanca - 5 linhas fixas, so a coluna DEFINICAO vira marcador
    gov = d.tables[0]
    tokens = ["{{GOV_CONVOCACAO}}", "{{GOV_PERIODICIDADE}}", "{{GOV_FORMATO}}", "{{GOV_PAUTA}}", "{{GOV_ATA}}"]
    esperado = ["Convoca", "Periodicidade", "Formato", "Pauta", "Registro"]
    for row, tok, esp in zip(gov.rows[1:6], tokens, esperado):
        if not row.cells[0].text.strip().startswith(esp):
            raise SystemExit("Linha inesperada na tabela de Governanca: " + row.cells[0].text)
        cell_token(row.cells[1], tok)

    # 5, 6, 7: tabelas repetidas por linha (1a linha de dados vira marcador, 2a e removida)
    specs = [
        (1, "ÁREA", ["{{AREA}}", "{{AREA_RESP}}"]),
        (2, "INPUT", ["{{INPUT}}", "{{INPUT_DESC}}", "{{INPUT_RESP}}"]),
        (3, "OUTPUT", ["{{OUTPUT}}", "{{OUTPUT_DESC}}", "{{OUTPUT_RESP}}"]),
    ]
    for idx, cab, toks in specs:
        t = d.tables[idx]
        if not t.rows[0].cells[0].text.strip().startswith(cab):
            raise SystemExit("Tabela %d inesperada: %s" % (idx + 1, t.rows[0].cells[0].text))
        for cell, tok in zip(t.rows[1].cells, toks):
            cell_token(cell, tok)
        for extra in list(t.rows)[2:]:
            extra._tr.getparent().remove(extra._tr)

    # titulos que vem logo depois de uma lista com marcador: sem isso o Word "cola" o titulo no ultimo item
    # (o estilo PargrafodaLista ignora o espaco entre paragrafos do mesmo estilo)
    for titulo in ("GOVERNANÇA E OPERACIONALIZAÇÃO", "CONTROLE DE EMISSÃO"):
        h = find(titulo)
        ppr = h._p.find(qn("w:pPr"))
        for old in ppr.findall(qn("w:contextualSpacing")):
            ppr.remove(old)
        cs = parse_xml('<w:contextualSpacing %s w:val="0"/>' % W)
        ref = None
        for tag in ("w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"):
            ref = ppr.find(qn(tag))
            if ref is not None:
                break
        if ref is not None:
            ref.addprevious(cs)
        else:
            ppr.append(cs)

    # capa: titulo (aparece na caixa de texto e na copia de compatibilidade)
    n = 0
    for p in d.element.body.iter(qn("w:p")):
        if ptext(p).strip() == "TÍTULO DE REGIMENTO":
            set_single_run(p, "{{TITULO}}")
            n += 1
    print("titulos da capa trocados:", n)

    d.save(dst)
    print("ok ->", dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
