# -*- coding: utf-8 -*-
"""Gera 'IT.GG.PRO.001 - MODELO IT com marcadores.docx' a partir do template oficial de IT.

Uso: python montar_modelo_it.py <template_IT.docx> <saida_com_marcadores.docx>

Marcadores colocados (so no corpo, igual ao modelo de POP):
  {{TITULO}}          titulo da capa
  {{OBJETIVO}}        texto da secao 1
  {{APLICACAO}}       texto da secao 2
  {{REFERENCIA_ITEM}} um paragrafo com marcador (bullet) por referencia (secao 3)
  {{SIGLA}} / {{SIGLA_DEF}}  linha repetida da tabela de definicoes (secao 4)
  {{ETAPAS}}          bloco da secao 5 (gerado pelo formulario, com imagens)
O Controle de Emissao e a capa/cabecalho/rodape ficam como no template (Processos preenche).
"""
import sys, copy
import docx
from docx.oxml.ns import qn
from docx.oxml import parse_xml

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
RPR = ('<w:rPr><w:rFonts w:ascii="Arial" w:eastAsia="Calibri" w:hAnsi="Arial" w:cs="Arial"/>'
       '<w:color w:val="000000" w:themeColor="text1"/><w:sz w:val="20"/><w:szCs w:val="20"/>'
       '<w:lang w:eastAsia="pt-BR"/></w:rPr>')


def ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def set_single_run(p_el, text):
    """Deixa o paragrafo com 1 so run (mantendo o rPr do primeiro) com o texto dado."""
    runs = p_el.findall(qn("w:r"))
    first = runs[0]
    for r in runs[1:]:
        p_el.remove(r)
    for child in list(first):
        if child.tag != qn("w:rPr"):
            first.remove(child)
    t = first.makeelement(qn("w:t"), {})
    t.text = text
    first.append(t)


def main(src, dst):
    d = docx.Document(src)
    body_ps = list(d.paragraphs)

    def find(startswith):
        for p in body_ps:
            if ptext(p._p).strip().startswith(startswith):
                return p
        raise SystemExit("Nao achei o paragrafo: " + startswith)

    # 1 e 2: textos simples
    set_single_run(find("Definir o objetivo")._p, "{{OBJETIVO}}")
    set_single_run(find("Descrever a Empresa")._p, "{{APLICACAO}}")

    # 3: referencias -> paragrafo com bullet, repetido por item
    ref = find("Relacionar as normas")
    new_p = parse_xml(
        '<w:p %s><w:pPr><w:pStyle w:val="PargrafodaLista"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr>'
        '<w:spacing w:after="0" w:line="360" w:lineRule="auto"/><w:jc w:val="both"/>%s</w:pPr>'
        '<w:r>%s<w:t>{{REFERENCIA_ITEM}}</w:t></w:r></w:p>' % (W, RPR, RPR))
    ref._p.addprevious(new_p)
    ref._p.getparent().remove(ref._p)

    # 4: remove o texto de instrucao e coloca marcadores na linha em branco da tabela
    instr = find("Identificar e descrever")
    instr._p.getparent().remove(instr._p)
    tbl = d.tables[0]
    row = tbl.rows[1]
    for cell, token in zip(row.cells, ("{{SIGLA}}", "{{SIGLA_DEF}}")):
        p = cell.paragraphs[0]._p
        for r in p.findall(qn("w:r")):
            p.remove(r)
        p.append(parse_xml('<w:r %s>%s<w:t>%s</w:t></w:r>' % (W, RPR, token)))

    # 5: troca os 3 paragrafos de instrucao por um unico marcador {{ETAPAS}}
    a = find("Inserir texto explicativo")
    b = find("Inserir imagem referente")
    c = find("Observação: Para cada passo")
    set_single_run(a._p, "{{ETAPAS}}")
    b._p.getparent().remove(b._p)
    c._p.getparent().remove(c._p)

    # capa: titulo (aparece na caixa de texto e na copia de compatibilidade)
    alvo = "TÍTULO DA INSTRUÇÃO DE TRABALHO"
    n = 0
    for p in d.element.body.iter(qn("w:p")):
        if ptext(p).strip() == alvo:
            set_single_run(p, "{{TITULO}}")
            n += 1
    print("titulos da capa trocados:", n)

    d.save(dst)
    print("ok ->", dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
